#!/usr/bin/env python
# -*-coding:utf-8 -*-
import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.utils.encoding import escape_uri_path


def export_excel(wb, filename):
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(content_type='application/vnd.ms-excel')
    response['Content-Disposition'] = 'attachment; filename={}'.format(
        escape_uri_path("{}.xlsx".format(filename))
    )
    response.write(output.getvalue())
    return response


def export_csv(rows, headers, filename):
    """
    导出 CSV 文件

    :param rows:     二维列表，每行一条记录
    :param headers:  列标题列表
    :param filename: 文件名（不带后缀）
    """
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename={}'.format(
        escape_uri_path("{}.csv".format(filename))
    )
    response.write(output.getvalue())
    return response


def import_csv(file_obj):
    """
    读取上传的 CSV 文件，返回 list[dict]

    :param file_obj: request.FILES['file']
    :return: list of OrderedDict
    """
    decoded = file_obj.read().decode('utf-8-sig')
    reader = csv.DictReader(StringIO(decoded))
    return list(reader)


def import_excel(file_obj, sheet_index=0):
    """
    读取上传的 Excel 文件，返回 list[dict]

    :param file_obj:     request.FILES['file']
    :param sheet_index:  工作表索引
    :return: list of dict
    """
    from openpyxl import load_workbook
    wb = load_workbook(file_obj, read_only=True)
    ws = wb.worksheets[sheet_index]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h else f'col_{i}' for i, h in enumerate(rows[0])]
    return [dict(zip(headers, row)) for row in rows[1:]]
